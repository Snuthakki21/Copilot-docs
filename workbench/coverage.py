"""Fail-closed, source-order accountability for the frozen process export.

Lines and semantic units have separate denominators. Verification is local,
source-derived and bounded; it never claims observed mainframe equivalence.
"""
import ast
import csv
import html
import json
import re
import sqlite3
from io import StringIO
from pathlib import Path
from openpyxl import Workbook, load_workbook
from .domain import encode, sha, safe_path, atomic_json, require

DISPOSITIONS = ('mapped_verified', 'mapped_unverified', 'blocked',
                'platform_replaced_verified', 'platform_replaced_unverified',
                'non_executable', 'out_of_scope')
VERIFIED = {'mapped_verified', 'platform_replaced_verified'}
COLUMNS = ('source_path', 'source_kind', 'source_hash', 'source_line', 'source_text',
           'source_line_hash', 'unit_id', 'source_start', 'source_end', 'disposition',
           'target_file', 'target_version', 'target_start', 'target_end', 'target_mappings',
           'tests', 'reason', 'replacement', 'evidence')


def _read(path):
    require(not path.is_symlink() and not any(p.is_symlink() for p in path.parents),
            'Evidence symlinks are not accepted')
    return path.read_bytes()


def _kind(path):
    suffix = Path(path).suffix.lower()
    return {'.cbl':'cobol_program', '.cob':'cobol_program', '.cobol':'cobol_program',
            '.cpy':'copybook', '.copy':'copybook', '.jcl':'jcl_job', '.bms':'bms_map',
            '.sql':'sql'}.get(suffix, 'other_source')


def _target(path, root, start, end):
    raw = _read(path)
    try: name = path.relative_to(root).as_posix()
    except ValueError: name = str(path)
    return {'file':name, 'version':sha(raw), 'start':start, 'end':end}


def _program_evidence(doc, root, base, name, p, global_errors, checkpoint=None):
    """Reproduce frozen expectations and compare complete stored target evidence."""
    from .fixtures import plan_cases, verify_program, adversarial_review
    from .target import emit_program
    state = {'verified':False, 'mappings':{}, 'tests':[], 'evidence':[], 'reason':'Verification is incomplete or unavailable.'}
    if p.get('blockers'):
        state['reason'] = 'Program contains unresolved source semantics; no complete target is credited.'
        return state
    try:
        if checkpoint:checkpoint()
        version = doc.get('program_versions', {}).get(name)
        require(isinstance(version,str) and re.fullmatch(r'[a-f0-9]{64}',version), 'Program target version is missing or malformed')
        target = root/'shared/target/python'/(version+'.py')
        raw = _read(target); code = raw.decode('utf-8')
        require(sha(raw) == version and code == emit_program(p), 'Target differs from its pinned generated version')
        function = ast.parse(code).body[0]
        for rule, node in zip(p['rules'], [n for n in function.body if isinstance(n, ast.If)]):
            state['mappings'][rule['id']] = _target(target, root, node.lineno, node.end_lineno)
        returns = [n for n in ast.walk(function) if isinstance(n, ast.Return)]
        state['mappings']['terminal'] = _target(target, root, returns[-1].lineno, returns[-1].end_lineno)
        # The actual execution adapter validates the source field layout before Python execution.
        from . import reference
        adapter = Path(reference.__file__).resolve()
        adapter_function = next(n for n in ast.parse(adapter.read_text()).body if isinstance(n, ast.FunctionDef) and n.name == 'input_errors')
        state['mappings']['layout'] = _target(adapter, root, adapter_function.lineno, adapter_function.end_lineno)
        require(not global_errors, 'Frozen process/evidence integrity failed; target replay and verified credit are prohibited')
        require(not doc.get('cancel_requested'), 'Cancelled: target execution prohibited; existing evidence is retained without new verification credit')
        require(doc.get('verification_finished'), 'Verification did not finish')
        require(doc.get('runs'), 'No synthetic run exists')
        run = doc['runs'][-1]; result = run.get('programs', {}).get(name)
        require(result is not None, 'Program is absent from the latest synthetic run')
        run_id = run['id']; identity_path = safe_path(base/'synthetic', run_id+'/'+name)
        expected_path = identity_path/'expected.json'; actual_path = identity_path/'actual-and-comparison.json'
        expected_raw = _read(expected_path); actual_raw = _read(actual_path)
        suite = json.loads(expected_raw)
        reproduced = plan_cases(p, doc.get('authorization', {}).get('seed', 21),
                                doc.get('authorization', {}).get('max_cases_per_program', 256))
        require(expected_raw == encode(reproduced), 'Frozen expected evidence changed or cannot be reproduced')
        require(actual_raw == encode(result), 'Stored actual evidence differs from its run record')
        replay = verify_program(p, code, suite, checkpoint=checkpoint)
        require(all(encode(result.get(k)) == encode(v) for k,v in replay.items()), 'Comparison evidence does not reproduce')
        require(len(suite['cases']) > 0 and result.get('expected_count') == len(suite['cases']), 'Synthetic run is empty or incomplete')
        require(not replay['differences'] and replay['coverage']['complete'], 'Comparison mismatch or missing modeled branch witness')
        adversarial=adversarial_review(p,code,suite,checkpoint=checkpoint)
        require(encode(result.get('adversarial')) == encode(adversarial) and adversarial.get('passed'), 'Adversarial evidence changed, is incomplete or contains unproved mutations')
        answers = (doc.get('answers') or {}).get('items', {})
        require(all(answers.get(r['id'], {}).get('answer') == 'Yes' and not answers.get(r['id'], {}).get('correction') for r in p['rules']),
                'SME confirmation is unresolved or has a correction')
        # Export the same frozen target version used by this run.
        copy_path = safe_path(base/'target', run_id+'/'+name+'.py')
        require(_read(copy_path) == raw, 'Run target copy is missing or changed')
        state['tests'] = [c['id'] for c in suite['cases']]
        state['evidence'] = [
            {'file':expected_path.relative_to(root).as_posix(), 'sha256':sha(expected_raw), 'kind':'frozen_source_expectations'},
            {'file':actual_path.relative_to(root).as_posix(), 'sha256':sha(actual_raw), 'kind':'reproduced_target_comparison'},
            {'file':copy_path.relative_to(root).as_posix(), 'sha256':sha(raw), 'kind':'executed_target_version'}]
        state['verified'] = not global_errors
        state['reason'] = 'SME-confirmed source IR, reproduced frozen synthetic expectations and matching target outputs; bounded local evidence.'
        if global_errors: state['reason'] = 'Source or analysis integrity failed; existing verification cannot receive credit.'
    except Exception as exc:
        state['reason'] = str(exc) or type(exc).__name__
        state['error'] = name+': '+state['reason']
    return state


def _job_evidence(doc, root, base, global_errors, checkpoint=None):
    from .target import emit_jobs
    from .orchestration import verify_jobs
    state = {'verified':False, 'mappings':{}, 'step_mappings':{}, 'tests':[], 'evidence':[], 'reason':'No verified job adapter evidence exists.'}
    if doc.get('cancel_requested'):
        state['reason']='Cancelled: target execution prohibited; existing job evidence is not reexecuted.'
        return state
    try:
        if checkpoint:checkpoint()
        require(not global_errors, 'Frozen process/evidence integrity failed; job target replay is prohibited')
        require(doc.get('verification_finished') and doc.get('runs'), 'Job verification is incomplete')
        run = doc['runs'][-1]['id']; folder = safe_path(base/'target', run)
        path = folder/'jobs.py'; raw = _read(path); code = raw.decode()
        require(code == emit_jobs(doc, doc.get('program_versions', {})), 'Job target is not the pinned generated orchestration')
        comparison_path = folder/'job-comparison.json'; comparison_raw = _read(comparison_path)
        result = verify_jobs(doc, root, code, checkpoint=checkpoint)
        require(comparison_raw == encode(result) and result.get('matched'), 'Job comparison is missing, changed or does not reproduce')
        for fn in ast.parse(code).body:
            if isinstance(fn, ast.FunctionDef) and fn.name.startswith('run_job_'):
                name = fn.name[len('run_job_'):].upper().replace('_','-')
                state['mappings'][name] = _target(path, root, fn.lineno, fn.end_lineno)
                job=next((j for j in doc['jobs'] if j['name'].upper()==name),None)
                if job:
                    for step,node in zip(job['steps'],[n for n in fn.body if isinstance(n,ast.If)]):
                        state['step_mappings'][(name,step['name'].upper())]=_target(path,root,node.lineno,node.end_lineno)
        order_answer=(doc.get('answers') or {}).get('items',{}).get('G_ORDER',{})
        require(order_answer.get('answer')=='Yes' and not order_answer.get('correction'),'Job order/dispatch SME confirmation is unresolved')
        state['verified'] = not global_errors
        state['tests'] = ['ordered_record_adapter_baseline']
        state['evidence'] = [{'file':comparison_path.relative_to(root).as_posix(), 'sha256':sha(comparison_raw), 'kind':'reproduced_job_comparison'}]
        state['reason'] = 'One reproduced ordered JSON-record adapter job baseline; native scheduler and dataset I/O are not reproduced.'
    except Exception as exc: state['reason'] = str(exc) or type(exc).__name__
    return state


def _manifest_integrity(doc,base):
    from .intake import parse_manifest
    pinned=doc.get('manifest_hash')
    require(isinstance(pinned,str) and re.fullmatch(r'[a-f0-9]{64}',pinned),
            'Frozen manifest baseline is missing; original intake must be recovered rather than repinned')
    raw=_read(base/'input/process-input.md')
    require(sha(raw)==pinned,'Frozen manifest bytes changed from the pinned intake baseline')
    manifest=parse_manifest(raw.decode('utf-8'))
    require(encode(manifest)==encode({key:doc.get(key) for key in ('id','name','jobs')}),
            'Frozen manifest identity/name/jobs differ from the pinned process document')


def _database_integrity(doc,root,base):
    if not doc.get('runs'):return
    run=doc['runs'][-1]
    path=safe_path(base/'target',run['id']+'/target.sqlite')
    relative=path.relative_to(base).as_posix()
    pinned=doc.get('artifact_hashes',{}).get(relative)
    require(isinstance(pinned,str) and re.fullmatch(r'[a-f0-9]{64}',pinned),
            'Frozen target database hash baseline is missing')
    raw=_read(path);require(sha(raw)==pinned,'Frozen target database bytes changed from the pinned result artifact')
    connection=sqlite3.connect(path.as_uri()+'?mode=ro',uri=True,timeout=1)
    try:
        tables=connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchall()
        require(tables==[('results',)],'Target database has unexpected or missing tables')
        schema=connection.execute('PRAGMA table_info(results)').fetchall()
        require([(row[1],row[2].upper(),row[5]) for row in schema]==[('program','TEXT',1),('case_id','TEXT',2),('result','TEXT',0)],
                'Target database result schema differs from the frozen execution contract')
        actual={(program,case_id):json.loads(result) for program,case_id,result in connection.execute('SELECT program,case_id,result FROM results')}
        expected={(name,row['case_id']):row['result'] for name,result in run.get('programs',{}).items() for row in result.get('actual',[])}
        require(encode(sorted((name,case_id,result) for (name,case_id),result in actual.items()))==encode(sorted((name,case_id,result) for (name,case_id),result in expected.items())),
                'Target database results are missing, unexpected, or differ from the frozen actual run')
        require(sha(_read(path))==pinned,'Target database changed during read-only inspection')
    finally:connection.close()


def build_coverage(doc, workspace_root, checkpoint=None):
    """Return one accounting row per original physical source line, in file order.

    Integrity failures produce explicit unverified/blocked rows and prevent
    completion. Inventory includes files excluded from the selected process.
    """
    root = Path(workspace_root).resolve(); base = safe_path(root, 'processes/'+doc['id'])
    analysis = doc.get('analysis') or {}; errors = []; sources = {}; source_ok = {}
    assets = {a['path']:a for a in analysis.get('assets', [])}
    programs = analysis.get('programs', {})
    for label,check in [('Manifest',lambda:_manifest_integrity(doc,base)),('Target database',lambda:_database_integrity(doc,root,base))]:
        try:check()
        except Exception as exc:errors.append(label+': '+(str(exc) or type(exc).__name__))
    for path, expected_hash in sorted(doc.get('source_files', {}).items()):
        frozen_text = assets.get(path, {}).get('source_text')
        if frozen_text is None:
            frozen_text = next((p.get('source_text') for p in programs.values() if p['path'] == path), None)
        try:
            raw = _read(safe_path(base/'input/sources', path)); text = raw.decode('utf-8')
            source_ok[path] = sha(raw) == expected_hash
            if not source_ok[path]: errors.append(path+': Frozen source hash changed')
            sources[path] = text if source_ok[path] or frozen_text is None else frozen_text
        except Exception as exc:
            source_ok[path] = False; sources[path] = frozen_text or ''
            errors.append(path+': Source snapshot unavailable: '+str(exc))
    try:
        require(bool(doc.get('source_files')), 'No source inventory exists')
        require(bool(analysis), 'Source analysis is not available')
        frozen = _read(base/'analysis/source-analysis.json')
        require(frozen == encode(analysis), 'Frozen source analysis changed or differs from process analysis')
        if all(source_ok.values()):
            from .source import analyze_sources
            require(encode(analyze_sources(sources, doc)) == frozen, 'Source analysis does not reproduce from frozen input')
    except Exception as exc: errors.append(str(exc) or type(exc).__name__)
    states = {name:_program_evidence(doc, root, base, name, p, errors, checkpoint) for name,p in programs.items()}
    integrity = list(errors)+[s['error'] for s in states.values() if s.get('error')]
    job_state = _job_evidence(doc, root, base, errors, checkpoint)
    by_path = {p['path']:(name,p) for name,p in programs.items()}
    referenced = {d['path'] for p in programs.values() for d in p.get('dependencies', [])}
    referenced.update(ref.rsplit(':',1)[0] for p in programs.values() for f in p.get('fields', {}).values() for ref in [f.get('source_ref','')] if ':' in ref)
    rows = []; inventory = []
    for path,text in sorted(sources.items()):
        asset = assets.get(path, {}); kind = asset.get('kind', _kind(path)); lines = text.splitlines()
        selected = path in by_path or path in referenced or (bool(asset) and asset.get('selected', True))
        if asset.get('scope_disposition') == 'out_of_scope': selected = False
        reason_scope = asset.get('scope_reason') or 'Exported asset is not referenced by the selected manifest programs, jobs or their dependencies.'
        inventory.append({'path':path, 'kind':kind, 'source_hash':doc['source_files'][path],
                          'physical_lines':len(lines), 'selected':selected, 'scope_reason':None if selected else reason_scope,
                          'integrity_verified':source_ok[path], 'source_text':text})
        program_entry = by_path.get(path)
        current_job=None
        copy_lines=[]
        if kind=='copybook':
            from .source import normalized_lines
            copy_lines=normalized_lines(text)
        for line_no,original in enumerate(lines,1):
            unit = 'line:'+path+':'+str(line_no); start = end = line_no
            disposition = 'blocked'; reason = 'Source behavior has no supported translation or verified replacement adapter.'
            replacement = None; mappings = []; tests = []; evidence = []
            stripped = original.strip(); state = None; mapping_key = None
            if kind=='jcl_job':
                declared=re.match(r'//([A-Z][A-Z0-9]{0,7})\s+JOB\b',stripped,re.I)
                if declared:current_job=declared[1].upper()
            if not selected:
                disposition = 'out_of_scope'; reason = reason_scope
            elif not source_ok[path]: reason = 'Frozen source integrity failed; original behavior cannot receive credit.'
            elif not stripped or (kind=='jcl_job' and original.startswith('//*')) or (kind=='sql' and stripped.startswith('--')) or (kind=='copybook' and copy_lines[line_no-1][0] in ('blank','comment')):
                disposition = 'non_executable'; reason = 'Blank line or source comment; preserved without executable conversion credit.'
            elif program_entry:
                name,p = program_entry; state = states[name]
                coverage = next((c for c in p.get('coverage', []) if c['line'] == line_no), {})
                original_disposition = coverage.get('disposition', 'unaccounted')
                if original_disposition in ('blank','comment','structure','paragraph'):
                    disposition = 'non_executable'; reason = 'Source comment, declaration header or label; no independent runtime behavior.'
                elif original_disposition == 'modeled':
                    rid = coverage.get('rule_id'); rule = next((r for r in p['rules'] if r['id'] == rid), None)
                    if rule:
                        start = rule['source_start']; end = rule['source_end']; unit = 'rule:'+path+':'+rid; mapping_key = rid
                elif original_disposition in ('data_layout','copybook'):
                    unit = 'layout:'+path+':'+str(line_no); mapping_key = 'layout'
                elif original_disposition == 'terminal':
                    unit = 'terminal:'+path+':'+str(line_no); mapping_key = 'terminal'
                    replacement = 'GOBACK/STOP RUN maps to the generated Python return adapter with record, trace and zero return code in the supported process boundary.'
                else:
                    matching = [b['message'] for b in p.get('blockers', []) if line_no in b.get('lines', [])]
                    reason = '; '.join(matching) or 'Source statement is unsupported or unaccounted; a semantic adapter is required.'
            elif kind == 'copybook':
                parents = [(name,p) for name,p in programs.items() if any(f.get('source_ref') == path+':'+str(line_no) for f in p.get('fields', {}).values())]
                if parents:
                    unit = 'layout:'+path+':'+str(line_no)
                    for name,p in parents:
                        s = states[name]; m = s['mappings'].get('layout')
                        if m: mappings.append(m)
                        tests.extend(name+':'+t for t in s['tests']); evidence.extend(s['evidence'])
                    disposition = 'mapped_verified' if all(states[n]['verified'] for n,_ in parents) else 'mapped_unverified'
                    reason = 'Source field layout is represented by the tested JSON-record input validation adapter; native record encoding is outside this mapping.'
                elif re.fullmatch(r'01\s+[A-Z][A-Z0-9-]*\.',copy_lines[line_no-1][1],re.I):
                    disposition = 'non_executable'; reason = 'Copybook group label; individual field mappings are accounted separately.'
                else: reason = 'Referenced copybook declaration has no supported field mapping.'
            elif kind == 'jcl_job':
                card = stripped.upper()
                if re.fullmatch(r'//[A-Z][A-Z0-9]{0,7}\s+EXEC\s+PGM=[A-Z][A-Z0-9-]*', card):
                    execute=re.fullmatch(r'//([A-Z][A-Z0-9]{0,7})\s+EXEC\s+PGM=([A-Z][A-Z0-9-]*)',card)
                    step_mapping=job_state['step_mappings'].get((current_job,execute[1]))
                    mappings=[step_mapping] if step_mapping else []
                    disposition = 'mapped_verified' if job_state['verified'] and mappings and all(s['verified'] for s in states.values()) else 'mapped_unverified'
                    tests = job_state['tests']; evidence = job_state['evidence']; reason = job_state['reason']
                elif ' DD ' in card: reason = 'Dataset allocation/read/write/disposition semantics require a concrete native-I/O replacement adapter; JSON-record tests do not reproduce DD behavior.'
                elif re.fullmatch(r"//([A-Z][A-Z0-9]{0,7})\s+JOB(?:\s+\([^)]*\)(?:,'[^']*')?)?",card):
                    job_name=re.match(r'//([A-Z][A-Z0-9]{0,7})',card)[1]
                    mapping=job_state['mappings'].get(job_name)
                    if mapping:mappings=[mapping]
                    replacement='Declared JCL job identity and ordered run-unit dispatch map to the named Python job function; source accounting metadata is preserved in the immutable source inventory.'
                    disposition='platform_replaced_verified' if job_state['verified'] and mappings and all(s['verified'] for s in states.values()) else 'platform_replaced_unverified'
                    tests=job_state['tests'];evidence=job_state['evidence']
                    reason=job_state['reason']+' This maps only the declared run unit; no external scheduler/submission/restart parity is claimed.'
                elif ' JOB' in card: reason = 'Unsupported native job scheduling/submission parameters have no verified replacement adapter.'
                else: reason = 'Unsupported JCL card or clause; no scheduler or I/O replacement is inferred.'
            elif kind == 'bms_map': reason = 'BMS/CICS screen behavior requires concrete action and screen mapping plus verification; no business UI replacement exists.'
            elif kind == 'sql': reason = 'Source SQL/database effects have no executable target mapping or verified database adapter.'
            if state and mapping_key:
                m = state['mappings'].get(mapping_key)
                if m: mappings = [m]
                tests = state['tests']; evidence = state['evidence']
                disposition = ('platform_replaced_' if replacement else 'mapped_')+('verified' if state['verified'] and mappings else 'unverified')
                reason = state['reason']
            if errors and disposition in VERIFIED:
                disposition = disposition.replace('_verified', '_unverified'); reason = 'Source or analysis integrity failed; mapping cannot receive verified credit.'
            target = mappings[0] if mappings else {}
            rows.append({'source_path':path, 'source_kind':kind, 'source_hash':doc['source_files'][path],
                         'source_line':line_no, 'source_text':original, 'source_line_hash':sha(original),
                         'unit_id':sha(doc['source_files'][path]+':'+unit), 'source_start':start, 'source_end':end,
                         'disposition':disposition, 'target_file':target.get('file'), 'target_version':target.get('version'),
                         'target_start':target.get('start'), 'target_end':target.get('end'), 'target_mappings':mappings,
                         'tests':sorted(set(tests)), 'reason':reason, 'replacement':replacement, 'evidence':evidence})
    counts = {d:sum(r['disposition'] == d for r in rows) for d in DISPOSITIONS}
    applicable = sum(r['disposition'] not in ('non_executable','out_of_scope') for r in rows)
    verified = sum(r['disposition'] in VERIFIED for r in rows)
    units = {}
    for row in rows: units.setdefault(row['unit_id'], []).append(row['disposition'])
    unit_applicable = [ds for ds in units.values() if any(d not in ('non_executable','out_of_scope') for d in ds)]
    unit_verified = sum(all(d in VERIFIED for d in ds) for ds in unit_applicable)
    unit_dispositions={d:0 for d in DISPOSITIONS}
    # A semantic unit is verified only when every line in its span is verified.
    for dispositions in units.values():
        category=next((d for d in ('blocked','platform_replaced_unverified','mapped_unverified','platform_replaced_verified','mapped_verified','non_executable','out_of_scope') if d in dispositions), 'blocked')
        unit_dispositions[category]+=1
    fully_accounted = bool(doc.get('source_files')) and len(inventory) == len(doc['source_files']) and all(source_ok.values()) and len(rows) == sum(f['physical_lines'] for f in inventory)
    summary = {'fully_accounted':fully_accounted, 'completion_eligible':fully_accounted and not integrity and bool(applicable) and applicable == verified and bool(doc.get('verification_finished')) and not doc.get('cancel_requested') and not [b for b in doc.get('blockers',[]) if b.get('kind')!='source_accountability'],
               'source_files':len(inventory), 'source_lines':len(rows), 'in_scope_files':sum(f['selected'] for f in inventory),
               'in_scope_lines':len(rows)-counts['out_of_scope'], 'out_of_scope_lines':counts['out_of_scope'],
               'non_executable_lines':counts['non_executable'], 'applicable_lines':applicable,
               'verified_applicable_lines':verified, 'line_verification_percent':round(100*verified/applicable,2) if applicable else None,
               'dispositions':counts, 'semantic_units':{'total':len(units), 'applicable':len(unit_applicable), 'verified':unit_verified, 'dispositions':unit_dispositions,
               'verification_percent':round(100*unit_verified/len(unit_applicable),2) if unit_applicable else None}, 'integrity_errors':integrity}
    return {'version':1, 'process_id':doc['id'], 'manifest_hash':doc.get('manifest_hash'), 'source_snapshot':analysis.get('source_snapshot'),
            'evidence_basis':'SOURCE_DERIVED_EXPECTED', 'observed_mainframe_parity':False,
            'denominators':{'source_lines':'All frozen exported physical lines, including non-executable and explicitly out-of-scope lines.',
                            'line_verification_percent':'Verified applicable in-scope physical lines / all applicable in-scope physical lines. Blocked and unverified lines remain in the denominator.',
                            'semantic_units':'Rules share one stable unit across their source span; layouts and unsupported lines remain distinct units. This is not a count of all possible legacy business rules.',
                            'known_rules':'Extracted known-rule SME verification is a separate metric, never total modernization percentage.'},
            'summary':summary, 'files':inventory, 'rows':rows}


def _cell(value):
    if value is None: return ''
    text = json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value,(list,dict)) else str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text


def write_coverage(model, report_root):
    root = Path(report_root); root.mkdir(parents=True, exist_ok=True)
    atomic_json(root/'coverage.json', model)
    output = StringIO(); writer = csv.writer(output); writer.writerow(COLUMNS)
    for row in model['rows']: writer.writerow([_cell(row.get(c)) for c in COLUMNS])
    (root/'coverage.csv').write_text(output.getvalue(), encoding='utf-8')
    book = Workbook(); summary = book.active; summary.title = 'Summary'; summary.append(['Measure','Value'])
    chunks=book.create_sheet('Text chunks');chunks.append(['Table','Row','Column','Chunk sequence','Text','Original SHA256','Encoding'])
    def workbook_cell(value,table,row_number,column):
        raw=json.dumps(value,ensure_ascii=False,sort_keys=True) if isinstance(value,(list,dict)) else '' if value is None else str(value)
        invalid=bool(re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',raw))
        preserved=json.dumps(raw,ensure_ascii=True) if invalid else raw
        if len(preserved)>32767 or invalid:
            for sequence,offset in enumerate(range(0,len(preserved),30000),1):
                chunks.append([table,row_number,column,sequence,preserved[offset:offset+30000],sha(raw),'json_string' if invalid else 'literal'])
                chunks.cell(chunks.max_row,5).data_type='s'
            return '[Full text in Text chunks: '+table+' row '+str(row_number)+' column '+column+']'
        return _cell(value)
    for key,value in model['summary'].items(): summary.append([key,workbook_cell(value,'Summary',summary.max_row+1,key)])
    sheet = book.create_sheet('Source lines'); sheet.append(list(COLUMNS))
    for row in model['rows']: sheet.append([workbook_cell(row.get(c),'Source lines',sheet.max_row+1,c) for c in COLUMNS])
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    for column in ('A','E','P','Q','R','S'): sheet.column_dimensions[column].width = 45
    inventory = book.create_sheet('Files'); file_columns = ('path','kind','source_hash','physical_lines','selected','scope_reason','integrity_verified','source_text')
    inventory.append(list(file_columns))
    for file in model['files']: inventory.append([workbook_cell(file.get(c),'Files',inventory.max_row+1,c) for c in file_columns])
    denominators = book.create_sheet('Denominators'); denominators.append(['Measure','Meaning'])
    for key,value in model['denominators'].items(): denominators.append([key,value])
    denominators.append(['Spreadsheet source text','Formula-leading CSV/XLSX values receive a protective apostrophe prefix. Oversized or XML-incompatible workbook values are preserved reversibly in Text chunks, with original hashes and literal/JSON encoding. JSON remains the exact canonical source text.'])
    book.save(root/'coverage.xlsx'); book.close()
    esc = lambda v: html.escape(json.dumps(v,ensure_ascii=False,sort_keys=True) if isinstance(v,(list,dict)) else '' if v is None else str(v), quote=True)
    body = ['<!doctype html><html lang="en"><meta charset="utf-8"><title>Source accountability</title>',
            '<style>body{font:14px system-ui;margin:24px}table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:6px;vertical-align:top}pre{white-space:pre-wrap}th{background:#edf2f7}</style>',
            '<h1>Source accountability</h1><p>Bounded source-derived verification. Observed mainframe parity is not established.</p>',
            '<pre>'+html.escape(json.dumps(model['summary'],indent=2))+'</pre>', '<h2>Denominators</h2><ul>']
    body.extend('<li>'+esc(k)+': '+esc(v)+'</li>' for k,v in model['denominators'].items())
    body.append('</ul><h2>Frozen file inventory</h2><table><tr>')
    body.extend('<th>'+esc(c)+'</th>' for c in file_columns[:-1]); body.append('</tr>')
    for file in model['files']: body.append('<tr>'+''.join('<td>'+esc(file.get(c))+'</td>' for c in file_columns[:-1])+'</tr>')
    body.append('</table><h2>Every source line, in original file order</h2><table><tr>')
    body.extend('<th>'+esc(c)+'</th>' for c in COLUMNS); body.append('</tr>')
    for row in model['rows']: body.append('<tr>'+''.join('<td><pre>'+esc(row.get(c))+'</pre></td>' for c in COLUMNS)+'</tr>')
    body.append('</table></html>'); (root/'coverage.html').write_text(''.join(body), encoding='utf-8')
    # Check meaningful row preservation and workbook reopenability before reports can certify output.
    require(json.loads((root/'coverage.json').read_text()) == model, 'Coverage JSON round-trip failed')
    require(len(list(csv.DictReader(StringIO((root/'coverage.csv').read_text())))) == len(model['rows']), 'Coverage CSV row count changed')
    reopened = load_workbook(root/'coverage.xlsx', read_only=True)
    try: require(reopened['Source lines'].max_row == len(model['rows'])+1, 'Coverage workbook row count changed')
    finally: reopened.close()
    return [root/name for name in ('coverage.json','coverage.csv','coverage.xlsx','coverage.html')]
