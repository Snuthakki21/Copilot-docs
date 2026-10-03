"""One editable, plain-language SME packet with immutable question identities."""
from io import BytesIO
from pathlib import Path
import html
import json
import tempfile
from .domain import require, sha, encode, decode, checked_zip, write_new, ValidationError

GLOBAL_QUESTIONS = [
 ('G_SCOPE','Does the listed program/job/step inventory cover this process?'),
 ('G_LAYOUT','Are the listed field names, lengths, numeric limits and record groups correct?'),
 ('G_ORDER','Are the listed job/step sequence, conditions, and input/output links correct?'),
 ('G_MATCH','Are the listed matching relationships and intentional missing-match cases interpreted correctly?'),
 ('G_TARGET','Is Python with SQLite/JSON files acceptable for this non-production POC?')]


def packet_document(process):
    a=process['analysis']
    items=[{'id':r['id'],'question':r['plain'],'evidence':'; '.join(r['source_refs']),'kind':'business_rule'} for r in a['rules']]
    context=json.dumps({'jobs':process['jobs'],'layouts':{n:p['fields'] for n,p in a['programs'].items()},'relationships':a['relationships'],'knowledge_input':process.get('knowledge_context')},ensure_ascii=False,indent=2)
    items += [{'id':k,'question':q,'evidence':'See process inventory/layout/relationship context','kind':'process_assumption'} for k,q in GLOBAL_QUESTIONS]
    if process.get('knowledge_context'):items.append({'id':'G_KNOWLEDGE','question':'Is the supplied background knowledge correct and applicable to this process? If No, describe the correction.','evidence':'See knowledge_input context and its SHA256','kind':'process_assumption'})
    suggestions=process.get('llm',{}).get('analysis',{})
    for n,text in enumerate(suggestions.get('questions',[])+suggestions.get('assumptions',[])):
        items.append({'id':f'LLM_{n:03d}','question':text,'evidence':'Unverified LLM suggestion; validate against source/context','kind':'provider_suggestion'})
    for i,b in enumerate(a['blockers']):items.append({'id':f'B_{i:03d}','question':'Is this unresolved item described correctly? If no, explain what it should do. '+b['message'],'evidence':b.get('path','Source/intake evidence'),'kind':'unresolved_item'})
    require(len(items)<=2000,'SME packet exceeds supported size; retain a scope blocker before issuing')
    doc={'version':1,'process_id':process['id'],'source_snapshot':a['source_snapshot'],'items':items,'context':context}
    doc['packet_hash']=sha(encode(doc))
    return doc


def export_packet(process, directory):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from docx import Document
    directory=Path(directory);document=packet_document(process)
    if directory.exists():
        existing=decode((directory/'packet.json').read_bytes())
        require(existing['packet_hash']==document['packet_hash'],'Existing packet differs; cannot issue a second packet')
        return existing
    directory.parent.mkdir(parents=True,exist_ok=True)
    temp=Path(tempfile.mkdtemp(prefix='.review-draft-',dir=directory.parent))
    book=Workbook();sheet=book.active;sheet.title='Checklist'
    sheet.append(['Item ID','What we understood','Source evidence','Category','Answer','If No, what should it be?','Reviewer'])
    def celltext(s):return "'"+s if s.lstrip().startswith(('=','+','-','@')) else s
    for item in document['items']:sheet.append([item['id'],celltext(item['question']),celltext(item['evidence']),item['kind'],'','',''])
    for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='18314F')
    for row in sheet.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
    for col,width in [('A',26),('B',80),('C',42),('D',22),('E',18),('F',60),('G',24)]:sheet.column_dimensions[col].width=width
    sheet.freeze_panes='E2';sheet.auto_filter.ref=sheet.dimensions
    dv=DataValidation(type='list',formula1='"Yes,No,Not sure"');dv.errorTitle='Choose an answer';dv.error='Use Yes, No or Not sure';dv.showErrorMessage=True;sheet.add_data_validation(dv);dv.add(f'E2:E{sheet.max_row}')
    meta=book.create_sheet('Metadata');meta.append(['Process ID',document['process_id']]);meta.append(['Packet hash',document['packet_hash']]);meta.append(['Source snapshot',document['source_snapshot']]);meta.append(['Evidence class','SOURCE_DERIVED_EXPECTED; no mainframe execution'])
    context=book.create_sheet('Context');context.append(['Process information']);context.column_dimensions['A'].width=120
    for i in range(0,len(document['context']),16000):
        context.append([document['context'][i:i+16000]]);context.cell(context.max_row,1).alignment=Alignment(wrap_text=True)
    book.save(temp/'sme-checklist.xlsx');book.close()
    word=Document();word.add_heading(process['name']+' — review checklist',0)
    word.add_paragraph('Check each statement. Answer Yes, No or Not sure in the Excel file. For No, write the correction. This is the only review round for this process. Missing or uncertain answers remain unresolved. Source-derived expectations are predictions, not observed mainframe results.')
    word.add_heading('Process context',1);word.add_paragraph(document['context'])
    for item in document['items']:
        word.add_heading(item['id'],2);word.add_paragraph(item['question']);word.add_paragraph('Evidence: '+item['evidence']);word.add_paragraph('Yes / No / Not sure. If No: __________________')
    word.save(temp/'sme-checklist.docx')
    rendered='<!doctype html><meta charset="utf-8"><title>SME checklist</title><h1>'+html.escape(process['name'])+'</h1><p>Return the Excel workbook. One review round. No mainframe execution.</p><pre>'+html.escape(document['context'])+'</pre>'+''.join('<section><h2>'+html.escape(i['id'])+'</h2><p>'+html.escape(i['question'])+'</p><small>'+html.escape(i['evidence'])+'</small></section>' for i in document['items'])
    write_new(temp/'sme-checklist.html',rendered.encode());write_new(temp/'packet.json',encode(document));temp.rename(directory)
    return document


def read_answers(data, packet, reviewer):
    from openpyxl import load_workbook
    require(isinstance(reviewer,str) and 0<len(reviewer.strip())<=160,'Name the reviewer responsible for this returned file')
    checked_zip(data).close();book=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        require({'Metadata','Checklist'}<=set(book.sheetnames),'Missing review sheets')
        meta=book['Metadata'];sheet=book['Checklist']
        require(meta['B1'].value==packet['process_id'] and meta['B2'].value==packet['packet_hash'] and meta['B3'].value==packet['source_snapshot'],'Packet identity/source version changed')
        require(sheet.max_row==len(packet['items'])+1 and sheet.max_column==7,'Checklist item set or dimensions changed')
        expected={x['id']:x for x in packet['items']};answers={}
        for cells in sheet.iter_rows(min_row=2):
            require(all(c.data_type!='f' for c in cells),'Formulas are not accepted in returned answers')
            values=[c.value for c in cells];rid,q,ref,kind,answer,correction,actor=values
            require(rid in expected and rid not in answers,'Unknown or duplicate question ID')
            original=expected[rid]
            def clean(s):return s[1:] if isinstance(s,str) and s.startswith("'") and s[1:].lstrip().startswith(('=','+','-','@')) else s
            require(clean(q)==original['question'] and clean(ref)==original['evidence'] and kind==original['kind'],'Original question/evidence changed')
            answer=str(answer or '').strip();correction=str(correction or '').strip();actor=str(actor or reviewer).strip()
            require(answer in ('Yes','No','Not sure',''),'Answer must be Yes, No, Not sure or blank')
            require(len(correction)<=4000 and 0<len(actor)<=160,'Returned text exceeds supported bounds')
            answers[rid]={'answer':answer or 'Unanswered','correction':correction,'reviewer':actor,'question':original['question'],'kind':kind}
        require(set(answers)==set(expected),'Missing checklist items')
        return {'packet_hash':packet['packet_hash'],'source_snapshot':packet['source_snapshot'],'reviewer':reviewer.strip(),'items':answers,'return_hash':sha(data)}
    finally:book.close()
