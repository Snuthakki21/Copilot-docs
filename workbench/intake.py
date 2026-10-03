"""Small, explicit Markdown/Excel process intake. Missing facts stay unknown."""
import re
from io import BytesIO
from .domain import ValidationError, identity, require, checked_zip

HEADERS = ['Job order', 'Job', 'Step order', 'Step', 'Program or utility', 'Input files/tables', 'Output files/tables', 'Condition or dependency']


def from_rows(pid, name, rows):
    identity(pid)
    require(isinstance(name, str) and 0 < len(name.strip()) <= 160, 'Process name is required and limited to 160 characters')
    require(0 < len(rows) <= 200, 'Supply between 1 and 200 job steps')
    jobs, seen, orders = {}, set(), {}
    for row in rows:
        require(len(row) == 8, 'Each job/step row must have eight columns')
        try: jo, so = int(row[0]), int(row[2])
        except (TypeError, ValueError) as exc: raise ValidationError('Job and step order must be positive integers') from exc
        require(0 < jo <= 1000 and 0 < so <= 1000, 'Job/step order out of range')
        job, step, program = map(lambda v: identity(str(v).strip()), (row[1], row[3], row[4]))
        require((job, so) not in seen, 'Duplicate step order within a job')
        require(not any(s['name'] == step for s in jobs.get(job, {}).get('steps', [])), 'Duplicate step name within a job')
        seen.add((job, so))
        require(jo not in orders or orders[jo] == job, 'Different jobs cannot share a job order')
        orders[jo] = job
        if job in jobs: require(jobs[job]['order'] == jo, 'One job has conflicting orders')
        jobs.setdefault(job, {'name': job, 'order': jo, 'steps': []})['steps'].append({
            'name': step, 'order': so, 'program': program,
            'inputs': [x.strip() for x in re.split('[;,]', str(row[5] or '')) if x.strip()],
            'outputs': [x.strip() for x in re.split('[;,]', str(row[6] or '')) if x.strip()],
            'condition': str(row[7] or 'Unknown').strip()})
    result = sorted(jobs.values(), key=lambda j: j['order'])
    for j in result: j['steps'].sort(key=lambda s: s['order'])
    return {'id': pid, 'name': name.strip(), 'jobs': result}


def parse_manifest(text):
    require(isinstance(text, str) and len(text) <= 128000, 'Markdown intake is too large')
    def attr(label):
        m = re.search(r'^\s*[-*]?\s*' + label + r'\s*:\s*`?([^`\n]+)`?\s*$', text, re.M | re.I)
        require(m is not None, f'Missing {label}')
        return m.group(1).strip()
    rows = []
    for line in text.splitlines():
        if not line.strip().startswith('|'): continue
        cells = [v.strip() for v in line.strip().strip('|').split('|')]
        if cells and cells[0].isdigit(): rows.append(cells)
    return from_rows(attr('Process ID'), attr('Process name'), rows)


def parse_intake_xlsx(data):
    from openpyxl import load_workbook
    checked_zip(data).close()
    book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        require('Intake' in book.sheetnames, 'Workbook needs an Intake sheet')
        ws = book['Intake']
        require(ws.max_row <= 205 and ws.max_column <= 8, 'Workbook exceeds intake bounds')
        for row in ws:
            require(all(c.data_type != 'f' for c in row), 'Formulas are not accepted in intake')
        require([ws.cell(4,c).value for c in range(1,9)] == HEADERS, 'Intake headers changed')
        rows = [[ws.cell(r,c).value for c in range(1,9)] for r in range(5,ws.max_row+1) if ws.cell(r,1).value is not None]
        return from_rows(str(ws['B1'].value or ''), str(ws['B2'].value or ''), rows)
    finally: book.close()
